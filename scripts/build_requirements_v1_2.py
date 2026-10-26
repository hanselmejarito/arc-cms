from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt

source = Path(r'C:\Users\hansel\Documents\ARC - Clinic Management System\ARC_Clinic_Management_System_Final_Business_System_Requirements_v1.1.docx')
output = Path(r'\\wsl$\Ubuntu\home\hplusplus\projects\clinic-codex\docs\ARC_Final_Business_and_System_Requirements_v1.2_For_Approval.docx')
doc = Document(source)

def replace(paragraph, before, after):
    if before not in paragraph.text:
        raise ValueError(f'Missing: {before}')
    for run in paragraph.runs:
        if before in run.text:
            run.text = run.text.replace(before, after)
            return
    # Source paragraphs for this document are single runs; this protects that layout.
    raise ValueError(f'Text split across runs: {before}')

def cell(table, row, column, before, after):
    replace(doc.tables[table].rows[row].cells[column].paragraphs[0], before, after)

def append_table_row(table, requirement_id, requirement, status):
    table = doc.tables[table]
    new_row = deepcopy(table.rows[-1]._tr)
    table._tbl.append(new_row)
    for index, value in enumerate((requirement_id, requirement, status)):
        new_cell = table.rows[-1].cells[index]
        new_cell.text = value

replace(doc.paragraphs[0], 'Version 1.1 - Updated Scope Baseline for Client Review', 'Version 1.2 - Scope Baseline for Client Approval')
replace(doc.paragraphs[0], 'Prepared from client meeting/transcriptions and supplied ARC/DOH forms.', 'Updated from client meetings, supplied ARC/DOH forms, and current system review on 26 October 2026.')
replace(doc.paragraphs[4], 'This document defines the current business and system requirements', 'This document defines the updated business and system requirements')
replace(doc.paragraphs[4], 'for client confirmation, solution design, database/ERD design, development planning, testing, and acceptance', 'for client approval, completion of remaining development, user acceptance testing, and deployment planning')
cell(0, 3, 0, 'Working Development Price', 'Development Fee')
cell(0, 4, 1, 'Define the agreed business scope, functional requirements, exclusions, and remaining configuration items before development.', 'Define the approved scope, exclusions, remaining configuration items, and UAT requirements for the current implementation.')
replace(doc.paragraphs[18], 'Owner/Admin users', 'Superadmin/Admin users')
replace(doc.paragraphs[117], 'Once this Version 1.1 document and the remaining configuration items are confirmed, it should be treated as the baseline scope for detailed system design and development.', 'Once this Version 1.2 document and the remaining configuration items are confirmed, it becomes the approved scope baseline for the current system and remaining development, testing and deployment work.')

cell(1, 4, 1, 'Owner/Admin', 'Superadmin/Admin')
cell(1, 5, 1, 'Roles discussed include Owner, Admin, Nurse/Staff, CPO/Officer, Secretary, and Cashier/operational staff.', 'The current system roles are Superadmin, Admin, Nurse, PhilHealth Officer and Secretary. Owner and Cashier are not separate application roles.')
cell(1, 5, 2, 'Role names may be refined', 'Current implemented role list; final operation permissions require UAT')
cell(2, 4, 2, 'numbering format to configure', 'current format ARC{branch code}-YY-0001; verify with the clinic')
cell(5, 6, 1, 'Record staff name/signature for each progress note.', 'Record the responsible staff name/signature for each progress note. Patient and visit registry are linked automatically; same-day recorded BP, temperature and animal status may prefill for clinician review. Current symptoms, assessment and plan require clinician entry.')
append_table_row(4, 'TRT-09', 'The current schedule is ID D0/D3/D7/D28, with D14 added when the animal is reported dead or lost within the 14-day observation window; IM D0/D3/D7/D14/D28; booster D0/D3; and PrEP D0/D7/D21.', 'Current implementation; clinic approval and UAT required')
append_table_row(4, 'TRT-10', 'More than seven days past the next due date marks the case Incomplete. Staff may still record a late dose; remaining due dates recalculate from its actual administration date.', 'Current implementation; clinic approval and UAT required')
cell(7, 12, 1, 'For Bioshoot, one vial has five injection-use units/sites; usage accumulates according to actual injections until the vial is fully consumed, at which point one vial is deducted.', 'Bioshoot, Speeda and Abhayrab each use five injection-use units/sites per vial. Usage accumulates by actual injection and one vial is deducted when fully consumed.')
cell(7, 14, 2, 'Confirmed', 'Required for UAT; automatic treatment deduction is not yet implemented')
cell(7, 15, 2, 'Confirmed', 'Required for UAT; automatic treatment deduction is not yet implemented')
append_table_row(7, 'INV-16', 'ChiroRab and VaxiRab N (PCEC) each use ten injection-use units/sites per vial. The open-vial counter displays uses out of ten; one vial is deducted when fully consumed.', 'Current consumption rule; verify against approved clinic protocol')
append_table_row(7, 'INV-17', 'ERIG (serum) stock is recorded through manual Stock Out according to actual vials used; it is not deducted automatically from the treatment entry.', 'Current workflow; verify actual usage with clinic staff')
append_table_row(8, 'SAL-11', 'Branch referral begins at a non-accredited branch and goes to a different PhilHealth-accredited branch. Acceptance creates a linked receiving registry and PhilHealth classification; an eligible earlier cash payment stays in history while a separate refund request is opened at the paying branch.', 'Current workflow; validate referral and refund cases in UAT')
cell(10, 2, 1, 'employee login is recorded as an attendance/time-in reference.', 'employee login and attendance are separate actions. Employees explicitly Time In and Time Out.')
cell(10, 2, 2, 'Confirmed concept', 'Current implemented workflow')
cell(11, 3, 1, 'including Excel/PDF where applicable', 'with print and XLSX export. PDF download remains an open requirement where applicable')
cell(11, 3, 2, 'Confirmed', 'Print and XLSX implemented; PDF download pending')
cell(13, 4, 1, 'Support PDF and Excel export for applicable reports/registries.', 'Support XLSX export and printing for applicable reports/registries. PDF download remains required where approved and must pass UAT before it is represented as delivered.')
cell(13, 4, 2, 'Confirmed', 'XLSX and print implemented; PDF download pending')
cell(15, 1, 1, 'current negotiated working scope', 'agreed development scope')
cell(15, 6, 1, 'Milestone-based payment schedule; total development fee remains PHP 120,000.', 'Milestone-based payment schedule totaling PHP 120,000. The first two milestones total 70% (PHP 84,000) and are due together after the final presentation and system demonstration.')
cell(15, 7, 1, 'due upon approval of the final scope and before development starts.', 'due after the final presentation and system demonstration, together with the Core Development Milestone, as agreed by the parties.')
cell(15, 8, 1, 'due upon completion/presentation of the agreed major/core modules.', 'due after the final presentation and system demonstration of the agreed major/core modules, together with the Downpayment / Project Start milestone.')

# The signature approves scope and commercial baseline; it does not certify UAT.
signature = doc.paragraphs[118]
for text in [
    'Approval of this document confirms the agreed business and system scope and the unchanged commercial amounts in Section 23. It is not confirmation that implementation, clinical validation, user acceptance testing, security checks or production deployment are complete.',
    'Before go-live, the parties will verify treatment schedules, inventory and payroll calculations, role and branch permissions, DOH report totals, production backup and recovery, and remaining PDF/settings work. Client-provided master data and current contribution rules in Section 21 still require confirmation. Payments already made, if any, are tracked separately from this scope document.'
]:
    paragraph = signature.insert_paragraph_before(text, style='Normal')
    paragraph.paragraph_format.space_after = Pt(7)

replace(next(p for p in doc.paragraphs if 'Client Representative:' in p.text), 'Client Representative:', 'Client Representative Name and Signature:')
replace(next(p for p in doc.paragraphs if 'Developer/Project Representative:' in p.text), 'Developer/Project Representative:', 'Developer/Project Representative Name and Signature:')
open_decision = next(p for p in doc.paragraphs if p.text.startswith('Final current DOH/clinic reporting templates'))
open_decision.insert_paragraph_before('Confirm the RER consent and signature method, including whether electronic signature capture or Save Draft/Finalize is required. These details are open design decisions until the client approves them.', style='List Bullet')

output.parent.mkdir(parents=True, exist_ok=True)
doc.save(output)
print(output)
